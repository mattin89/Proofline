"""Build benign, deterministic PDF and DOCX fixtures for upload-parser smoke tests."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parent
BODY = (
    "NeuroWeave Robotics is a fictional parser-validation company developing a "
    "soft robotic hand rehabilitation device. The authorized plan reports a paid "
    "pilot hypothesis, a hospital procurement path, and a prototype validation plan."
)
BOUNDARY = (
    "This fixture contains no real founder, customer, revenue, clinical, or private data. "
    "Every statement is synthetic and must remain founder-provided evidence."
)


def build_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading = document.styles["Heading 1"]
    heading.font.name = "Calibri"
    heading.font.size = Pt(16)
    heading.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)

    document.core_properties.title = "Proofline Upload Parser Smoke Test"
    document.core_properties.author = "Proofline test fixture"
    document.add_heading("Proofline Upload Parser Smoke Test", level=1)
    document.add_paragraph(BODY)
    document.add_paragraph(BOUNDARY)
    document.save(path)


def build_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    heading = ParagraphStyle(
        "ProoflineHeading",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=19,
        textColor="#2E74B5",
        spaceBefore=16,
        spaceAfter=8,
    )
    body = ParagraphStyle(
        "ProoflineBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=11,
        leading=13.2,
        spaceAfter=6,
    )
    document = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title="Proofline Upload Parser Smoke Test",
        author="Proofline test fixture",
    )
    document.build([
        Paragraph("Proofline Upload Parser Smoke Test", heading),
        Spacer(1, 4),
        Paragraph(BODY, body),
        Paragraph(BOUNDARY, body),
    ])


if __name__ == "__main__":
    build_docx(ROOT / "business-plan-smoke.docx")
    build_pdf(ROOT / "business-plan-smoke.pdf")
